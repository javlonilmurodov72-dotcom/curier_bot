import os
import json
import glob
import asyncio
from datetime import datetime, timedelta
import openpyxl

from aiogram import Bot, Dispatcher, F, types
from aiogram.enums import ContentType
from aiogram.filters import CommandStart, Command
from aiogram.types import FSInputFile, InlineKeyboardMarkup, InlineKeyboardButton
from groq import Groq
from openai import OpenAI
from aiohttp import web

# ==========================================
# API KALITLAR VA ADMIN SOZLAMALARI (RENDER ENV)
# ==========================================
BOT_TOKEN = os.environ.get("BOT_TOKEN")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
OPENROUTER_API_KEY = os.environ.get("OPENROUTER_API_KEY")

# ADMIN TELEGRAM ID RAQAMLARI (Masalan: "123456789,987654321")
ADMIN_IDS_RAW = os.environ.get("ADMIN_ID", "123456789")
ADMIN_IDS = [int(x.strip()) for x in ADMIN_IDS_RAW.split(",") if x.strip().isdigit()]

groq_client = Groq(api_key=GROQ_API_KEY)
qwen_client = OpenAI(
    base_url="https://openrouter.ai/api/v1",
    api_key=OPENROUTER_API_KEY,
)

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

# ------------------------------------------
# Render Web Service Port Binding (Dummy Server)
# ------------------------------------------
async def handle_ping(request):
    return web.Response(text="Bot is running smoothly 24/7!")

async def start_dummy_web_server():
    app = web.Application()
    app.router.add_get('/', handle_ping)
    runner = web.AppRunner(app)
    await runner.setup()
    port = int(os.environ.get("PORT", 8080))
    site = web.TCPSite(runner, '0.0.0.0', port)
    await site.start()
    print(f"Soxta veb-server port {port} da ishga tushdi.")

# ------------------------------------------
# 1. Soat 04:00 chegara va 3 kunlik tozalash
# ------------------------------------------
def get_working_date() -> datetime:
    """Soat 04:00 gacha bo'lgan vaqtni o'tgan kunga tegishli deb hisoblaydi."""
    now = datetime.now()
    if now.hour < 4:
        return now - timedelta(days=1)
    return now

def get_today_excel_file() -> str:
    work_date = get_working_date().strftime("%Y-%m-%d")
    return f"report_{work_date}.xlsx"

def auto_clean_old_excel_files():
    """3 kundan (72 soat) o'tgan Excel fayllarni avtomatik o'chirish"""
    now = datetime.now()
    for file_path in glob.glob("report_*.xlsx"):
        try:
            date_part = file_path.replace("report_", "").replace(".xlsx", "")
            file_date = datetime.strptime(date_part, "%Y-%m-%d")
            if now - file_date > timedelta(days=3):
                os.remove(file_path)
                print(f"Eski fayl o'chirildi: {file_path}")
        except Exception as e:
            print(f"Fayl o'chirishda xatolik: {e}")

# ------------------------------------------
# 2. Buyruqlar va Sanalar bilan Excel yuklash
# ------------------------------------------
@dp.message(CommandStart())
async def start_handler(message: types.Message):
    welcome_text = (
        "👋 **Xush kelibsiz!**\n\n"
        "Mashina raqamingiz hamda ombordan chiqayotganingiz yoki borayotgan ofisingiz haqida **dumaloq video** yuboring.\n"
        "AI ma'lumotlarni smena va mashina raqami bo'yicha Excel'ga yozib boradi.\n\n"
        "📊 Excel hisobotni yuklab olish: /excel (Faqat adminlar uchun)"
    )
    await message.answer(welcome_text, parse_mode="Markdown")

@dp.message(Command("excel"))
async def send_excel_report_menu(message: types.Message):
    if message.from_user.id not in ADMIN_IDS:
        await message.answer("❌ **Ruxsat berilmadi!**\nExcel hisobotni faqat **Adminlar** yuklab olishi mumkin.")
        return

    auto_clean_old_excel_files()
    files = sorted(glob.glob("report_*.xlsx"), reverse=True)

    if not files:
        await message.answer("⚠️ Hali hech qanday hisobot fayllari mavjud emas.")
        return

    keyboard_buttons = []
    for file_path in files:
        date_str = file_path.replace("report_", "").replace(".xlsx", "")
        btn = InlineKeyboardButton(
            text=f"📅 {date_str} hisoboti", 
            callback_data=f"download_{date_str}"
        )
        keyboard_buttons.append([btn])

    reply_markup = InlineKeyboardMarkup(inline_keyboard=keyboard_buttons)
    await message.answer("📊 **Qaysi sanadagi Excel hisobotni yuklab olmoqchisiz?**", reply_markup=reply_markup, parse_mode="Markdown")

@dp.callback_query(F.data.startswith("download_"))
async def handle_excel_download(callback: types.CallbackQuery):
    if callback.from_user.id not in ADMIN_IDS:
        await callback.answer("❌ Sizga ruxsat berilmagan!", show_alert=True)
        return

    date_str = callback.data.replace("download_", "")
    file_path = f"report_{date_str}.xlsx"

    if os.path.exists(file_path):
        await callback.answer("⏳ Fayl yuklanmoqda...")
        doc = FSInputFile(file_path, filename=f"Kuryerlar_Hisoboti_{date_str}.xlsx")
        await callback.message.answer_document(doc, caption=f"📊 **{date_str}** kungi kuryerlar va mashinalar hisoboti.")
    else:
        await callback.answer("❌ Fayl topilmadi yoki o'chirilgan!", show_alert=True)

# ------------------------------------------
# 3. Groq Whisper orqali ovozni matnga o'girish
# ------------------------------------------
def speech_to_text(video_path: str) -> str:
    with open(video_path, "rb") as audio_file:
        transcription = groq_client.audio.transcriptions.create(
            file=(video_path, audio_file.read()),
            model="whisper-large-v3",
            language="uz",
            response_format="text"
        )
    return transcription

# ------------------------------------------
# 4. Qwen AI orqali matnni tahlil qilish
# ------------------------------------------
def parse_info_with_qwen(text: str) -> dict:
    prompt = f"""
    Siz kuryerlar yuborgan audio matnlarini tahlil qiluvchi yordamchisiz.
    Quyidagi matn kuryer tomonidan aytilgan:
    "{text}"

    Matndan quyidagi ma'lumotlarni ajratib oling:
    1. "car_number": Kuryer aytgan mashina raqami (masalan: "01A777AA" yoki "777", topilmasa "Noma'lum").
    2. "is_departure": Kuryer ombor/sklad/ofisdan YO'LGA CHIQAN bo'lsa `true`, agar biror ofisga YETIB BORGAN bo'lsa `false`.
    3. "location_name": Kuryer borgan ofis/PVZ yoki manzil nomi (masalan: "Chilonzor ofisi", "Sklad", "Yunusobod").

    Javobni FAQAT QUYIDAGI JSON formatida qaytaring:
    {{
      "car_number": "01A777AA",
      "is_departure": false,
      "location_name": "Chilonzor ofisi"
    }}
    """

    try:
        response = qwen_client.chat.completions.create(
            model="qwen/qwen-2.5-72b-instruct",
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"}
        )
        content = response.choices[0].message.content
        return json.loads(content)
    except Exception as e:
        print(f"Parsing Xatosi: {e}")
        return {"car_number": "Noma'lum", "is_departure": False, "location_name": "Noma'lum manzil"}

# ------------------------------------------
# 5. Smena va dinamik ustunlar bo'yicha Excel'ga yozish
# ------------------------------------------
def write_to_smena_excel(car_number: str, is_departure: bool, location_name: str, raw_text: str):
    auto_clean_old_excel_files()
    excel_file = get_today_excel_file()

    headers = ["Sana", "Smena", "Mashina raqami", "Ombordan chiqish vaqti", "1-Manzil / Ofis", "2-Manzil / Ofis", "3-Manzil / Ofis"]

    if not os.path.exists(excel_file):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        wb.save(excel_file)

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active

    now = datetime.now()
    work_date_str = get_working_date().strftime("%Y-%m-%d")
    time_str = now.strftime("%H:%M")

    if 4 <= now.hour < 14:
        smena_name = "1-Smena (Kunduzgi)"
    else:
        smena_name = "2-Smena (Kechki)"

    target_row = None
    for row in range(2, ws.max_row + 1):
        cell_smena = str(ws.cell(row=row, column=2).value or "")
        cell_car = str(ws.cell(row=row, column=3).value or "")
        if smena_name in cell_smena and car_number != "Noma'lum" and car_number in cell_car:
            target_row = row
            break

    if not target_row:
        target_row = ws.max_row + 1
        ws.cell(row=target_row, column=1, value=work_date_str)
        ws.cell(row=target_row, column=2, value=smena_name)
        ws.cell(row=target_row, column=3, value=car_number)

    if is_departure:
        ws.cell(row=target_row, column=4, value=f"{location_name} ({time_str})")
    else:
        col_idx = 5
        while ws.cell(row=target_row, column=col_idx).value is not None:
            col_idx += 1
        
        header_cell = ws.cell(row=1, column=col_idx).value
        if not header_cell:
            ws.cell(row=1, column=col_idx, value=f"{col_idx-4}-Manzil / Ofis")

        ws.cell(row=target_row, column=col_idx, value=f"{location_name} ({time_str})")

    wb.save(excel_file)

# ------------------------------------------
# 6. Dumaloq videolarni qabul qilish
# ------------------------------------------
@dp.message(F.content_type == ContentType.VIDEO_NOTE)
async def handle_video_note(message: types.Message):
    status_msg = await message.reply("⏳ Video qabul qilindi. AI tahlil qilmoqda...")

    video_note = message.video_note
    file_info = await bot.get_file(video_note.file_id)
    video_path = f"video_{message.message_id}.mp4"
    
    await bot.download_file(file_info.file_path, destination=video_path)

    try:
        raw_text = speech_to_text(video_path)
        parsed = parse_info_with_qwen(raw_text)

        car_num = parsed.get("car_number", "Noma'lum")
        is_departure = parsed.get("is_departure", False)
        location_name = parsed.get("location_name", "Ofis")

        if car_num == "Noma'lum":
            car_num = message.from_user.full_name or "Kuryer"

        write_to_smena_excel(car_num, is_departure, location_name, raw_text)

        res_text = (
            f"✅ **Excel'ga saqlandi!**\n\n"
            f"🚗 **Mashina raqami:** `{car_num}`\n"
            f"📍 **Joylashuv:** {location_name}\n"
            f"⏱ **Harakat:** {'Ombordan chiqdi' if is_departure else 'Ofisga yetib keldi'}\n"
            f"🎙 **Eshitilgan matn:** _{raw_text}_"
        )

        await status_msg.edit_text(res_text, parse_mode="Markdown")

    except Exception as e:
        await status_msg.edit_text(f"❌ Xatolik yuz berdi: {str(e)}")

    finally:
        if os.path.exists(video_path):
            os.remove(video_path)

async def main():
    print("Mavjud kunlik Excel fayllarni tanlab yuklovchi Bot ishga tushdi...")
    # Render portini faollashtirish
    await start_dummy_web_server()
    # Bot polling
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
